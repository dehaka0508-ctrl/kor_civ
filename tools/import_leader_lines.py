"""지도자 대사 엑셀 → korciv/data/leader_lines.json (+ 작성 기준 → docs/지도자_대사_작성기준.md).

사용: python3 tools/import_leader_lines.py 지도자대사.xlsx   (openpyxl 필요)
엑셀 '지도자 대사' 시트 열: No, 카테고리, 지도자, 언어, 첫 조우, 우호 선언, 전쟁 발발, 플레이어 비난,
동맹 체결, 휴전, 플레이어에게 멸망할 때, 플레이어를 멸망시켰을 때.
대사가 바뀌어 새 한자가 생기면 tools/build_cjk_font.py 로 보조 글꼴을 다시 만든다.
"""
import json
import os
import sys

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
from korciv.leaders import LEADERS  # noqa: E402

KINDS = ("meet", "friend", "war", "denounce", "alliance", "peace", "defeated", "victory")
ALIAS = {"선왕": "발해 선왕"}


def main(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    rows = list(wb["지도자 대사"].iter_rows(values_only=True))
    by_name = {l["name"]: l["key"] for l in LEADERS}
    out, unknown = {}, []
    for r in rows[1:]:
        if not r or not r[2]:
            continue
        name = ALIAS.get(str(r[2]).strip(), str(r[2]).strip())
        key = by_name.get(name)
        if key is None:
            unknown.append(name)
            continue
        out[key] = {"name": name, "lang": str(r[3] or "").strip()}
        for k, v in zip(KINDS, r[4:12]):
            out[key][k] = str(v or "").strip()
    missing = [l["name"] for l in LEADERS if l["key"] != "cus" and l["key"] not in out]
    with open(os.path.join(ROOT, "korciv", "data", "leader_lines.json"), "w", encoding="utf-8") as f:
        json.dump(out, f, ensure_ascii=False, indent=1)
    if "작성 기준" in wb.sheetnames:
        lines = ["# 지도자 대사 작성 기준", "", "대사 원본은 엑셀에서 `tools/import_leader_lines.py`로 "
                 "`korciv/data/leader_lines.json`에 옮긴다.", "", "| 항목 | 기준 |", "|---|---|"]
        for a, b in wb["작성 기준"].iter_rows(values_only=True):
            if a:
                lines.append(f"| {a} | {str(b or '').replace('|', '/')} |")
        with open(os.path.join(ROOT, "docs", "지도자_대사_작성기준.md"), "w", encoding="utf-8") as f:
            f.write("\n".join(lines) + "\n")
    print(f"{len(out)}명 저장", "· 이름을 못 찾음: " + ", ".join(unknown) if unknown else "",
          "· 대사 없음: " + ", ".join(missing) if missing else "")


if __name__ == "__main__":
    main(sys.argv[1])
